from datetime import time

import openpyxl

from paralisacao.preparar import gravar_planilha_padrao, montar_linhas
from paralisacao.resumo import ler_resumo


def _resumo_minimo(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B5"] = "Navio"
    ws["B7"] = "TI SUNRISE"
    ws["Q5"] = "Agente"
    ws["R5"] = "Terminal"
    ws["Q7"] = "MARINAV"
    ws["J19"] = "Produto"
    ws["J20"] = "NÍQUEL"
    ws["C31"] = "Paralisações - 16/10/26"
    ws["O35"] = "Motivo da paralisação"
    ws["C37"] = time(13, 0)
    ws["K37"] = time(13, 15)
    ws["O37"] = "DDS"
    ws["AH37"] = "Troca de Turno"
    ws["C38"] = time(0, 45)
    ws["K38"] = time(1, 0)
    ws["O38"] = "CHUVA"
    ws["C39"] = time(1, 0)
    ws["K39"] = time(1, 15)
    ws["O39"] = "DDS"
    ws["C44"] = "Paralisações - 17/10/26"
    ws["C48"] = time(7, 0)
    ws["K48"] = time(7, 15)
    ws["O48"] = "DDS"
    wb.save(path)


class TestLerResumo:
    def test_extrai_capa_e_eventos(self, tmp_path):
        caminho = tmp_path / "resumo.xlsx"
        _resumo_minimo(caminho)

        resumo = ler_resumo(caminho)

        assert resumo.navio == "TI SUNRISE"
        assert resumo.produto == "NÍQUEL"
        assert resumo.agencia == "MARINAV"
        assert [(e.data, e.inicio, e.fim, e.motivo) for e in resumo.eventos] == [
            ("16/10/2026", "13:00", "13:15", "DDS"),
            ("16/10/2026", "00:45", "01:00", "CHUVA"),
            ("16/10/2026", "01:00", "01:15", "DDS"),
            ("17/10/2026", "07:00", "07:15", "DDS"),
        ]


class TestMontarLinhas:
    def test_turno_usa_a_data_do_bloco(self, tmp_path):
        caminho = tmp_path / "resumo.xlsx"
        _resumo_minimo(caminho)
        linhas = montar_linhas(ler_resumo(caminho))

        assert [(ln.registro, ln.data, ln.periodo, ln.inicio) for ln in linhas] == [
            ("1", "16/10/2026", "13:00/19:00", "13:00"),
            ("2", "16/10/2026", "19:00/01:00", "00:45"),
            ("3", "16/10/2026", "01:00/07:00", "01:00"),
            ("4", "17/10/2026", "07:00/13:00", "07:00"),
        ]

    def test_gravacao_volta_pelo_leitor_atual(self, tmp_path):
        from paralisacao.planilha import ler_planilha

        caminho = tmp_path / "resumo.xlsx"
        saida = tmp_path / "padrao.xlsx"
        _resumo_minimo(caminho)
        resumo = ler_resumo(caminho)
        gravar_planilha_padrao(
            saida,
            resumo,
            montar_linhas(resumo),
            programacao="1766",
            operador="30002",
        )

        dados = ler_planilha(str(saida))
        assert dados.prog_navio == "1766"
        assert len(dados.capas) == 4
        assert dados.capas[1].periodo == "19:00/01:00"
        assert dados.capas[1].paralisacoes[0].motivo == "CHUVA"
