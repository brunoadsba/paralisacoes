from datetime import time
from pathlib import Path

import openpyxl
import pytest

from paralisacao.config import Config
import paralisacao.mapeamento as mapeamento
from paralisacao.mapeamento import codigo_motivo, salvar_codigo_motivo
from paralisacao.preparar import montar_linhas, pendencias
from paralisacao.resumo import EventoResumo, ResumoEmbarque
from paralisacao.ui.ambiente import url_operacao
from paralisacao.ui.app import app
from paralisacao.ui.estado import LOCK, STATE

RESUMO_REAL = Path(r"C:\Users\danyp\Downloads\TI SUNRISE - Resumo de embarque.xlsx")


def _config(tmp_path) -> Config:
    (tmp_path / "data").mkdir()
    (tmp_path / "data" / "motivos.json").write_text(
        '{"mapeamentos": {"DDS": "CONVENIENCIA DO USUARIO", "CHUVA": "CHUVA"}, "codigos": {}}',
        encoding="utf-8",
    )
    (tmp_path / "data" / "tipo-paralizacoes.md").write_text(
        "| 1 | 5 | CHUVA |\n| 2 | 34 | CONVENIENCIA DO USUARIO |\n",
        encoding="utf-8",
    )
    return Config(
        url="https://exemplo/",
        login="",
        senha="",
        planilha="",
        headless=False,
        primeiro_plano=False,
        timeout_ms=1000,
        slow_mo_ms=0,
        pausa_visual_ms=0,
        manter_aberto_seg=0,
        estadia_label="",
        operador="30002",
        project_root=tmp_path,
    )


def _resumo(path: Path) -> None:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["B5"] = "Navio"
    ws["B7"] = "TI SUNRISE"
    ws["Q5"] = "Agente"
    ws["Q7"] = "MARINAV"
    ws["J19"] = "Produto"
    ws["J20"] = "NÍQUEL"
    ws["C31"] = "Paralisações - 16/10/26"
    ws["O35"] = "Motivo da paralisação"
    ws["C37"] = time(13, 0)
    ws["K37"] = time(13, 15)
    ws["O37"] = "DDS"
    wb.save(path)


@pytest.fixture
def client():
    with LOCK:
        STATE.update({
            "usuario": "",
            "senha": "",
            "modo": "qas",
            "resumo": None,
            "linhas": [],
            "programacao": "",
            "planilha": "",
            "operador": "30002",
            "log": [],
            "phase": "Aguardando início",
            "running": False,
            "uma_ok": False,
            "proc": 0,
            "tot": 0,
            "situacao": {},
            "relatorio": "",
        })
    app.config["TESTING"] = True
    return app.test_client()


def test_url_qas_sem_default_aspx():
    url = url_operacao("qas")
    assert url == "https://ilheus.codeba.gov.br/qas/openportcodeba/"
    assert "default.aspx" not in url


def test_login_mostra_qas(client):
    resposta = client.get("/login")
    assert resposta.status_code == 200
    texto = resposta.get_data(as_text=True)
    assert "Paralisações" in texto
    assert "Testes (QAS)" in texto


def test_lote_permanece_bloqueado(client):
    resposta = client.post("/executar", data={"escopo": "lote"})
    assert "primeira paralisação" in resposta.get_data(as_text=True)
    assert b"disabled" in client.get("/execucao").data


def test_codigo_novo_entra_no_json(tmp_path):
    config = _config(tmp_path)
    anterior = mapeamento._CODIGOS_MOTIVO_CACHE
    mapeamento._CODIGOS_MOTIVO_CACHE = {}
    try:
        salvar_codigo_motivo(config, "Troca de Turno", 434)
        assert codigo_motivo("TROCA DE TURNO", config) == 434
        linhas = montar_linhas(ResumoEmbarque(
            navio="N", produto="P", agencia="A",
            eventos=[EventoResumo("16/10/2026", "13:00", "13:15", "Troca de Turno")],
        ))
        assert pendencias(linhas, config, "1766") == []
    finally:
        mapeamento._CODIGOS_MOTIVO_CACHE = anterior


def test_resumo_minimo_no_painel(client, tmp_path):
    caminho = tmp_path / "resumo.xlsx"
    _resumo(caminho)
    with caminho.open("rb") as arquivo:
        resposta = client.post(
            "/resumo",
            data={"resumo": (arquivo, "resumo.xlsx")},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    texto = resposta.get_data(as_text=True)
    assert "TI SUNRISE" in texto
    assert "DDS" in texto


@pytest.mark.skipif(not RESUMO_REAL.exists(), reason="resumo real fora do repositório")
def test_sunrise_doze_paralisacoes_e_cinco_capas(client):
    with RESUMO_REAL.open("rb") as arquivo:
        resposta = client.post(
            "/resumo",
            data={"resumo": (arquivo, RESUMO_REAL.name)},
            content_type="multipart/form-data",
            follow_redirects=True,
        )
    texto = resposta.get_data(as_text=True)
    assert "5 capas" in texto
    assert "12 paralisações" in texto
    assert "TI SUNRISE" in texto
