import re
from dataclasses import dataclass, field
from datetime import datetime, time
from pathlib import Path

import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

from paralisacao.mapeamento import normalizar

_DATA_BLOCO = re.compile(r"(\d{1,2})/(\d{1,2})/(\d{2,4})")


@dataclass
class EventoResumo:
    data: str
    inicio: str
    fim: str
    motivo: str


@dataclass
class ResumoEmbarque:
    navio: str
    produto: str
    agencia: str
    eventos: list[EventoResumo] = field(default_factory=list)


def ler_resumo(caminho: str | Path) -> ResumoEmbarque:
    """Lê o resumo de embarque da Intermarítima e devolve navio, produto e paralisações."""
    wb = openpyxl.load_workbook(caminho, data_only=True)
    ws = _aba_paralisações(wb)
    eventos = _eventos(ws)
    if not eventos:
        raise ValueError("Nenhuma paralisação encontrada no resumo de embarque.")
    return ResumoEmbarque(
        navio=_valor_rotulo(ws, "NAVIO"),
        produto=_valor_rotulo(ws, "PRODUTO"),
        agencia=_valor_rotulo(ws, "AGENTE"),
        eventos=eventos,
    )


def _aba_paralisações(wb) -> Worksheet:
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            if _data_de_titulo(row):
                return ws
    raise ValueError("Aba com blocos de paralisações não encontrada no resumo.")


def _valor_rotulo(ws: Worksheet, rotulo: str) -> str:
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None or normalizar(str(cell.value)) != rotulo:
                continue
            for linha in range(cell.row + 1, cell.row + 6):
                texto = _texto(ws.cell(linha, cell.column).value)
                if texto:
                    return texto
    return ""


def _eventos(ws: Worksheet) -> list[EventoResumo]:
    eventos: list[EventoResumo] = []
    data_bloco = ""
    coluna_motivo = 0
    for row in ws.iter_rows():
        titulo = _data_de_titulo(row)
        if titulo:
            data_bloco = titulo
            continue
        coluna = _coluna_motivo(row)
        if coluna:
            coluna_motivo = coluna
            continue
        if not data_bloco or not coluna_motivo:
            continue
        evento = _evento_da_linha(row, data_bloco, coluna_motivo)
        if evento:
            eventos.append(evento)
    return eventos


def _data_de_titulo(row) -> str:
    for cell in row:
        texto = _texto(cell.value)
        if "PARALIS" not in normalizar(texto):
            continue
        encontrado = _DATA_BLOCO.search(texto)
        if not encontrado:
            continue
        dia, mes, ano = (int(parte) for parte in encontrado.groups())
        if ano < 100:
            ano += 2000
        return f"{dia:02d}/{mes:02d}/{ano:04d}"
    return ""


def _coluna_motivo(row) -> int:
    for cell in row:
        if cell.value and "MOTIVO" in normalizar(str(cell.value)):
            return cell.column
    return 0


def _evento_da_linha(row, data_bloco: str, coluna_motivo: int) -> EventoResumo | None:
    tempos: list[str] = []
    for cell in row:
        valor = cell.value
        if isinstance(valor, datetime):
            tempos.append(f"{valor.hour:02d}:{valor.minute:02d}")
        elif isinstance(valor, time):
            tempos.append(f"{valor.hour:02d}:{valor.minute:02d}")
    motivo = ""
    for cell in row:
        if cell.column == coluna_motivo:
            motivo = _texto(cell.value)
            break
    if len(tempos) < 2 or not motivo or normalizar(motivo).startswith("TOTAL"):
        return None
    return EventoResumo(data_bloco, tempos[0], tempos[1], motivo)


def _texto(valor) -> str:
    if valor is None:
        return ""
    return str(valor).strip()
